"""Councils in Scotland, Wales and Northern Ireland: net revenue spending by service, from each nation's official statistics.

Each nation collects and classifies council spending differently, so the nations are compared only within themselves, never
with each other or with England. Per-resident rates use the ONS mid-year population estimates for every council.

Scotland: Scottish Government, Scottish Local Government Finance Statistics (SLGFS), "LA level - net revenue expenditure and
income by service and type" (£ thousand, one year per publication).
Wales: Welsh Government, StatsWales "Revenue outturn expenditure by authority and service" (£ thousand, 2001-02 onward).
"""
import csv
import re
import statistics
import urllib.parse

import analysis as A

SCOTLAND_SERVICES = [
    ("edu", "Education Services", "Education", "age0_17"), ("cul", "Culture and Related Services", "Culture and related services", "pop"),
    ("soc", "Social Work Services", "Social work", "pop"), ("trans", "Roads and Transport", "Roads and transport", "pop"),
    ("env", "Environmental Services", "Environmental services", "pop"), ("plan", "Building, Planning and Development", "Building, planning and development", "pop"),
    ("cen", "Central Services", "Central services", "pop"), ("hous", "Non-HRA Housing", "Housing (excluding council housing account)", "pop"),
    ("trad", "Trading Services", "Trading services", "pop"),
]
WALES_SERVICES = [
    ("edu", "Education", "Education", "age0_17"), ("soc", "Social services", "Social services", "pop"),
    ("trans", "Roads and transport", "Roads and transport", "pop"), ("env", "Local environmental services", "Local environmental services", "pop"),
    ("cul", "Libraries, culture, heritage, sport and recreation", "Libraries, culture, heritage, sport and recreation", "pop"),
    ("plan", "Planning and economic development", "Planning and economic development", "pop"),
    ("hous", "Council fund housing and housing benefit", "Council fund housing and housing benefit", "pop"),
    ("law", "Law and order and protective services", "Law and order and protective services (older years)", "pop"),
    ("ctb", "Council tax benefit and administration", "Council tax benefit and administration", "pop"),
    ("oth", "Other revenue expenditure", "Other revenue expenditure", "pop"),
]
WALES_TOTAL = "Revenue expenditure"
SCOTLAND_NATION = {"S92000003": "Scotland", "W92000004": "Wales", "N92000002": "Northern Ireland"}


def norm(name):
    return re.sub(r"[^a-z0-9]", "", str(name).lower().replace("&", "and"))


def pop_by_name(pop, prefix):
    return {norm(v["name"]): (k, v) for k, v in pop.items() if k.startswith(prefix)}


# ------------------------------------------------------------------ Scotland
def discover_scotland(get_text, today=None, doc_pattern="LA Level - Net Revenue Expenditure and Income by Service and Type"):
    """Find the newest SLGFS 'LA level' workbook by trying each financial year's publication page, newest first."""
    import datetime
    t = today or datetime.date.today()
    y = t.year if t.month >= 4 else t.year - 1
    for k in range(4):
        start = y - k
        slug = f"scottish-local-government-finance-statistics-{start}-{str(start + 1)[2:]}"
        try:
            html = get_text(f"https://www.gov.scot/publications/{slug}/documents/")
        except Exception:
            continue
        for m in re.finditer(r'href="([^"]+\.xlsx)"', html):
            href = urllib.parse.unquote(m.group(1)).replace("+", " ")
            if doc_pattern.lower() in href.lower():
                u = m.group(1)
                return (u if u.startswith("http") else "https://www.gov.scot" + u), f"{start}-{str(start + 1)[2:]}"
    raise RuntimeError("no SLGFS LA-level workbook found on gov.scot")


def scotland_year(path):
    """The financial year the workbook covers, from the title of its 'Scotland' sheet ('... Income in 2024-25 by Service and Type')."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Scotland"]
    ws.reset_dimensions()
    first = next(ws.iter_rows(values_only=True, max_row=1))[0]
    m = re.search(r"\b(20\d\d-\d\d)\b", str(first))
    if not m:
        raise RuntimeError(f"cannot read the year from {first!r}")
    return m.group(1)


def read_scotland(path, pop):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    by = pop_by_name(pop, "S12")
    out, total_row = {}, None
    for sn in wb.sheetnames:
        ws = wb[sn]
        ws.reset_dimensions()
        rows = list(ws.iter_rows(values_only=True))
        hi = next((i for i, r in enumerate(rows) if r and r[0] == "Expenditure / Income Type"), None)
        if hi is None:
            continue
        head = [str(x) if x is not None else "" for x in rows[hi]]
        net = next((r for r in rows[hi + 1:] if r and r[0] == "Net Revenue Expenditure"), None)
        if net is None:
            continue
        vals = {h: net[i] for i, h in enumerate(head) if i and isinstance(net[i], (int, float))}
        if sn == "Councils":
            total_row = vals
            continue
        hit = by.get(norm(sn))
        if hit:
            code, p = hit
            out[code] = {"code": code, "name": p["name"], "v": {k: vals.get(col) for k, col, _, _ in SCOTLAND_SERVICES}, "gf": vals.get("Total General Fund"), "hra": vals.get("HRA"), "tot": vals.get("Total General Fund + HRA")}
    return out, total_row


# --------------------------------------------------------------------- Wales
def read_wales(path, pop):
    by = pop_by_name(pop, "W06")
    keys = {s: k for k, s, _, _ in WALES_SERVICES}
    out, totals, per_head = {}, {}, {}
    for r in csv.DictReader(open(path, encoding="utf-8-sig")):
        try:
            v = float(r["Data values"])
        except ValueError:
            continue
        a, y, s, d = r["Authority"], r["Year"], r["Service"], r["Data description"]
        if a == "Total Unitary Authorities":
            if d == "£ thousand":
                totals.setdefault(y, {})[keys.get(s, "tot" if s == WALES_TOTAL else s)] = v
            continue
        hit = by.get(norm(a))
        if not hit:
            continue
        code, p = hit
        c = out.setdefault(code, {"code": code, "name": p["name"], "y": {}})
        if d == "£ thousand":
            c["y"].setdefault(y, {})[keys.get(s, "tot" if s == WALES_TOTAL else s)] = v
        elif d == "£ per head" and s == WALES_TOTAL:
            per_head.setdefault(code, {})[y] = v
    return out, totals, per_head


# ----------------------------------------------------------------- analysis
def robust(xs):
    return A.robust_center_spread(xs)


def peer_stats(councils, services, year_of):
    """{service: {n, median, p10, p90}} of £ per unit across the nation's councils."""
    out = {}
    for key, _, _, denom in services:
        xs = []
        for c in councils:
            v = year_of(c).get(key)
            d = c["pop"].get(denom) if c.get("pop") else None
            if v is not None and d:
                xs.append(v * 1000 / d)
        if len(xs) >= 8:
            s = sorted(xs)
            out[key] = {"n": len(xs), "median": statistics.median(xs), "p10": s[int(.1 * (len(s) - 1))], "p90": s[int(.9 * (len(s) - 1))]}
    return out


def attach_pop(councils, pop):
    for c in councils:
        p = pop.get(c["code"])
        c["pop"] = {"pop": round(p["pop"]), "age0_17": round(p["age0_17"]), "age65": round(p["age65"])} if p else None


def build_scotland(path, pop, year):
    cs, total_row = read_scotland(path, pop)
    councils = sorted(cs.values(), key=lambda c: c["name"])
    attach_pop(councils, pop)
    stats = peer_stats(councils, SCOTLAND_SERVICES, lambda c: c["v"])
    checks = scotland_checks(cs, total_row, pop)
    return {
        "id": "scotland", "label": "Scotland", "year": year, "years": [year], "unit": "£ thousand", "single_year": True,
        "services": [{"key": k, "label": lab, "denom": d} for k, _, lab, d in SCOTLAND_SERVICES],
        "councils": [{"code": c["code"], "name": c["name"], "pop": c["pop"], "v": {k: v for k, v in c["v"].items()}, "gf": c["gf"], "hra": c["hra"], "tot": c["tot"]} for c in councils],
        "stats": [{"key": k, **{a: round(b, 2) if isinstance(b, float) else b for a, b in s.items()}} for k, s in stats.items()],
        "peer_note": "All 32 councils are unitary authorities, so each is compared with the others. The three island councils (Orkney, Shetland, Na h-Eileanan Siar) are very sparsely populated and cost more per resident for most services.",
        "notes": [
            "Figures are net revenue expenditure by service (gross service expenditure minus service income), general fund, from councils' audited accounts where available, as defined by the Scottish Government's Local Financial Return.",
            "Social work spending is after income from Integration Joint Boards (IJBs), through which councils and NHS boards jointly run adult social care, so it is not comparable with England or Wales.",
            "Housing is non-HRA housing only (housing benefit administration and homelessness); council housing is paid for from the separate housing revenue account (HRA), shown separately.",
            "Trading services (for example leisure or building services run commercially) can be negative because they earn more than they cost.",
        ],
    }, checks


def build_wales(path, pop, real_factor):
    cs, totals, per_head = read_wales(path, pop)
    councils = sorted(cs.values(), key=lambda c: c["name"])
    attach_pop(councils, pop)
    years = sorted({y for c in councils for y in c["y"]})
    latest = years[-1]
    stats = peer_stats(councils, WALES_SERVICES, lambda c: c["y"].get(latest, {}))
    checks = wales_checks(cs, totals, per_head, pop, latest)
    return {
        "id": "wales", "label": "Wales", "year": latest, "years": years, "unit": "£ thousand", "single_year": False,
        "services": [{"key": k, "label": lab, "denom": d} for k, _, lab, d in WALES_SERVICES],
        "councils": [{"code": c["code"], "name": c["name"], "pop": c["pop"], "y": c["y"]} for c in councils],
        "totals": totals,
        "stats": [{"key": k, **{a: round(b, 2) if isinstance(b, float) else b for a, b in s.items()}} for k, s in stats.items()],
        "peer_note": "All 22 councils are unitary authorities, so each is compared with the others.",
        "notes": [
            "Figures are the Welsh Government's revenue outturn expenditure by service for the 22 unitary councils (what they actually spent), from the Revenue Outturn (RO) return, in £ thousand. Fire, police and national park authorities are not included.",
            "Education is spent on schools and related services run by the council; it is shown per child aged 0-17.",
            "Housing and housing benefit and council tax benefit are large because they include benefits paid out and later recovered from central government.",
        ],
    }, checks


# -------------------------------------------------------------------- checks
def scotland_checks(cs, total_row, pop):
    cks = []
    gap = []
    for c in cs.values():
        gap.append((c["name"], c["gf"], sum((c["v"].get(k) or 0) for k, *_ in SCOTLAND_SERVICES)))
    cks.append(A._check("dev_sc_services", "Scotland: each council's nine service lines add up to its total general fund net revenue expenditure (£ thousand)", gap, 5))
    if total_row:
        rows = [(col, total_row.get(col) or 0, sum((c["v"].get(k) or 0) for c in cs.values())) for k, col, *_ in SCOTLAND_SERVICES]
        rows += [("Total general fund", total_row.get("Total General Fund") or 0, sum(c["gf"] or 0 for c in cs.values())), ("Housing revenue account", total_row.get("HRA") or 0, sum(c["hra"] or 0 for c in cs.values())),
                 ("Total general fund + HRA", total_row.get("Total General Fund + HRA") or 0, sum(c["tot"] or 0 for c in cs.values()))]
        cks.append(A._check("dev_sc_councils", "Scotland: the 32 council returns add up to the published all-councils line, by service (£ thousand)", rows, 5))
    cks.append(A._check("dev_sc_count", "Scotland: all 32 council areas found with population (count)", [("council areas", 32, len(cs))], 0))
    cks.append(_pop_check(pop, "S12", "S92000003", "Scotland"))
    return _tag(cks)


def wales_checks(cs, totals, per_head, pop, latest):
    cks = []
    keys = [k for k, *_ in WALES_SERVICES]
    gap = []
    for c in cs.values():
        for y, v in c["y"].items():
            if "tot" in v:
                gap.append((f"{c['name']} {y}", v["tot"], sum(v.get(k) or 0 for k in keys)))
    cks.append(A._check("dev_wa_services", "Wales: each council's service lines add up to its total revenue expenditure, every year (£ thousand)", gap, 5))
    rows = []
    for y, t in totals.items():
        s = sum((c["y"].get(y, {}).get("tot") or 0) for c in cs.values())
        if t.get("tot") is not None and s:
            rows.append((y, t["tot"], s))
    cks.append(A._check("dev_wa_councils", "Wales: the 22 council returns add up to the published unitary authority total, every year (£ thousand)", rows, 5))
    # our per-head (ONS mid-2024) against the Welsh Government's own per-head figure (its own population base)
    ph = []
    for code, c in cs.items():
        t, p = c["y"].get(latest, {}).get("tot"), pop.get(code)
        theirs = per_head.get(code, {}).get(latest)
        if t and p and theirs:
            ph.append((c["name"], theirs, t * 1000 / p["pop"]))
    ck = A._check("dev_wa_perhead", f"Wales: £ per resident here (ONS mid-2024 population) against the Welsh Government's own £ per head, {latest}", ph, 60,
                  "The two use different population estimates, so small gaps are expected; a large gap would point to a mistake.")
    cks.append(ck)
    cks.append(_pop_check(pop, "W06", "W92000004", "Wales"))
    return _tag(cks)


def _pop_check(pop, prefix, nation_code, label):
    s = sum(p["pop"] for k, p in pop.items() if k.startswith(prefix))
    return A._check(f"dev_pop_{prefix.lower()}", f"{label}: council area populations add up to the ONS {label} total (people)", [("mid-2024", pop[nation_code]["pop"], s)], 1)


def _tag(cks):
    for c in cks:
        c["source"] = "Local councils (Scotland, Wales, Northern Ireland)"
        if len(c["rows"]) > 10:
            c["rows"] = sorted(c["rows"], key=lambda r: -abs(r["gap"]))[:10]
            c["note"] = (c["note"] + " " if c["note"] else "") + "Showing the 10 largest gaps."
    return cks
