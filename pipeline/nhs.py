"""NHS England referral-to-treatment (RTT) waiting times against health spending.

The monthly national series (April 2007 onwards) is one spreadsheet, "RTT Overview Timeseries Including Estimates
for Missing Trusts". NHS England's statistics pages are fronted by a bot-protection challenge, but the same pages
are published through the site's public WordPress REST API (wp-json), which serves the page text and file links
normally; the pipeline reads the link to the newest file from there, so the series refreshes without any manual step.
"""
import datetime
import json
import re

LINK_TEXT = "Overview Timeseries Including Estimates for Missing Trusts"
BASE = "https://www.england.nhs.uk/statistics/wp-json/wp/v2/pages?slug="
COLS = {"median": 3, "p92": 4, "within18_pct": 8, "over18": 10, "over52": 12, "over65": 16, "over78": 18, "over104": 20, "waiting": 22,
        "completed": 41, "new_periods": 38}
THRESHOLD_18W = 92.0        # the constitutional standard: 92% of waiting patients within 18 weeks


def fy_slugs(today):
    """Newest first: rtt-data-2026-27, rtt-data-2025-26 ..."""
    y = today.year if today.month >= 4 else today.year - 1
    return [f"rtt-data-{y - k}-{str(y - k + 1)[2:]}" for k in range(3)]


def find_file_link(page_html):
    for m in re.finditer(r'href="([^"]+\.xlsx?)"[^>]*>([^<]+)<', page_html):
        if LINK_TEXT.lower() in m.group(2).lower().replace(" ", " "):
            return m.group(1)
    return None


def discover(get_json, today=None):
    """get_json(url) -> parsed JSON. Returns the newest time-series file link."""
    for slug in fy_slugs(today or datetime.date.today()):
        pages = get_json(BASE + slug)
        if pages:
            link = find_file_link(pages[0]["content"]["rendered"])
            if link:
                return link
    raise RuntimeError("no RTT time series link found through the NHS England REST API")


def num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


def parse_rows(rows):
    """rows: worksheet rows (tuples). Returns {'YYYY-MM': {metric: value}} for months with a waiting-list figure."""
    out = {}
    for r in rows:
        d = r[2] if len(r) > 2 else None
        if not isinstance(d, datetime.datetime):
            continue
        rec = {k: num(r[i]) if len(r) > i else None for k, i in COLS.items()}
        if rec["waiting"] is None:
            continue
        for k in ("within18_pct",):
            if rec[k] is not None:
                rec[k] = rec[k] * 100
        out[d.strftime("%Y-%m")] = rec
    return out


def read(path):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    ws.reset_dimensions()
    return parse_rows(list(ws.iter_rows(values_only=True)))


def fy_of(month):
    y, m = int(month[:4]), int(month[5:])
    y0 = y if m >= 4 else y - 1
    return f"{y0}-{str(y0 + 1)[2:]}"


def build(series, fn_years, health_real, health_gdp):
    months = sorted(series)
    fy_end = {}
    for mth in months:
        if mth.endswith("-03"):
            fy_end[fy_of(mth)] = mth
    spend, gdp = dict(zip(fn_years, health_real)), dict(zip(fn_years, health_gdp))
    years = [y for y in sorted(fy_end) if y in spend]
    ann = {y: series[fy_end[y]] for y in years}
    latest = months[-1]
    peak = max(months, key=lambda m: series[m]["waiting"])
    return {
        "months": months,
        "waiting_m": [round(series[m]["waiting"] / 1e6, 3) for m in months],
        "within18_pct": [None if series[m]["within18_pct"] is None else round(series[m]["within18_pct"], 1) for m in months],
        "over52": [series[m]["over52"] for m in months],
        "median": [series[m]["median"] for m in months],
        "years": years,
        "spend_real_bn": {y: round(spend[y], 2) for y in years},
        "spend_gdp": {y: gdp.get(y) for y in years},
        "year_end": {y: fy_end[y] for y in years},
        "waiting_year_end_m": {y: round(ann[y]["waiting"] / 1e6, 2) for y in years},
        "within18_year_end": {y: None if ann[y]["within18_pct"] is None else round(ann[y]["within18_pct"], 1) for y in years},
        "latest": latest, "peak": peak, "standard_18w": THRESHOLD_18W,
        "notes": [
            "Spending is the Treasury's Health function for the whole UK (in 2025-26 prices); the waiting list is for England only. Health spending in England is about 85% of the UK total and the two have moved together, but this is a trend indicator rather than a like-for-like ratio.",
            "The waiting list counts pathways (referrals), not people: the same patient can be on it more than once. The figures use NHS England's estimates for trusts that did not report (England-wide totals).",
            "Waiting times are for consultant-led treatment only. The list is affected by activity in primary care, social care (delayed discharges), staffing and strikes, as well as by spending.",
            "From February 2024 community service pathways stopped being counted, which lowered the list slightly; earlier months are not restated.",
        ],
    }
