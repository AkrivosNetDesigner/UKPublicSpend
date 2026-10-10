"""Parsers for the HM Treasury PESA / Public Spending Statistics workbooks."""
import re
import openpyxl

YEAR_RE = re.compile(r"^\s*(\d{4}-\d{2})")
FOOTNOTE_RE = re.compile(r"\s*(\(\d+\)[,\s]*)+$")


def num(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        return None


def clean_label(s):
    s = str(s).replace("\n", " ").strip()
    s = FOOTNOTE_RE.sub("", s).strip()
    return re.sub(r"\s+", " ", s)


def rows_of(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [tuple(r) for r in wb[sheet].iter_rows(values_only=True)]
    finally:
        wb.close()


def year_header(rows, start=0):
    """Index + [(col, year)] of the first row at/after `start` with >=3 year labels."""
    for i in range(start, len(rows)):
        cols = [(c, YEAR_RE.match(str(v)).group(1)) for c, v in enumerate(rows[i])
                if v is not None and YEAR_RE.match(str(v))]
        if len(cols) >= 3:
            return i, cols
    return None, []


def parse_wide(path, sheet, header_start=0, label_col=0):
    """Generic label x year table. Returns (years, [{label, level, values}]).

    Stops at the first blank label or a footnote row such as "(1) ...".
    """
    rows = rows_of(path, sheet)
    hi, cols = year_header(rows, header_start)
    years = [y for _, y in cols]
    out, parent = [], None
    for r in rows[hi + 1:]:
        lab = r[label_col] if len(r) > label_col else None
        if lab is None or str(lab).strip() == "":
            if out:
                break
            continue
        if re.match(r"^\s*\(\d+\)", str(lab)) or str(lab).startswith("http"):
            break
        label = clean_label(lab)
        child = label.lower().startswith("of which")
        if child:
            label = re.sub(r"(?i)^of which:\s*", "", label)
        else:
            parent = label
        out.append({"label": label, "level": 1 if child else 0,
                    "parent": parent if child else None,
                    "values": [num(r[c]) for c, _ in cols]})
    return years, out


def parse_aggregates(path):
    """Table 4.1: TME, current spending, net investment 1985-86 -> 2025-26."""
    rows = rows_of(path, "4_1")
    out = []
    for r in rows:
        m = YEAR_RE.match(str(r[1])) if len(r) > 1 and r[1] else None
        if not m:
            continue
        out.append({
            "year": m.group(1),
            "current": num(r[2]), "current_real": num(r[3]), "current_gdp": num(r[4]),
            "depreciation": num(r[6]),
            "netinv": num(r[8]), "netinv_real": num(r[9]), "netinv_gdp": num(r[10]),
            "tme": num(r[12]), "tme_real": num(r[13]), "tme_gdp": num(r[14]),
        })
    return out


def parse_table_1_1(path):
    """Table 1.1 Total Managed Expenditure, £ million. Returns (years, lines)."""
    rows = rows_of(path, "Table_1_1")
    hi, cols = year_header(rows)
    years = [y for _, y in cols]
    lines, section = [], ""
    for r in rows[hi + 1:]:
        lab = r[0]
        if lab is None or str(lab).strip() == "":
            continue
        if re.match(r"^\s*\(\d+\)", str(lab)):
            break
        vals = [num(r[c]) for c, _ in cols]
        label = clean_label(lab)
        if all(v is None for v in vals):
            section = label
            continue
        lines.append({"label": label, "section": section, "values": vals})
    return years, lines


def parse_region_block(path, sheet, header_start=0):
    """Label x year block. Some tables repeat the year labels for a second set of
    columns (e.g. per-cent shares); only the first set is returned in `values`,
    the second goes to `extra`."""
    rows = rows_of(path, sheet)
    hi, cols = year_header(rows, header_start)
    seen, first, second = set(), [], []
    for c, y in cols:
        if y in seen or second:
            second.append((c, y))
        else:
            seen.add(y)
            first.append((c, y))
    cols = first
    years = [y for _, y in cols]
    out = []
    for r in rows[hi + 1:]:
        lab = r[0]
        if lab is None or str(lab).strip() == "":
            if out:
                break
            continue
        if re.match(r"^\s*\(\d+\)", str(lab)):
            break
        row = {"label": clean_label(lab), "values": [num(r[c]) for c, _ in cols]}
        if second:
            row["extra"] = [num(r[c]) for c, _ in second]
        out.append(row)
    return years, out, hi


def parse_database(path):
    """PSS database -> list of dict rows. Amounts are £ thousand (accruals)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Database"]
    it = ws.iter_rows(values_only=True)
    hdr = next(it)
    idx = {h: i for i, h in enumerate(hdr)}
    out = []
    for r in it:
        if not r[idx["Year"]]:
            continue
        out.append({
            "year": r[idx["Year"]],
            "econ": r[idx["Economic Budget Code"]],
            "ring": r[idx["Economic Ringfence Code"]],
            "ctrl": r[idx["Control Budget Code"]],
            "detail": r[idx["Control Budget Detail Code"]],
            "dept": r[idx["SR Group Long Name"]],
            "code": r[idx["Segment Code"]],
            "name": r[idx["Segment Long Name"]],
            "group": r[idx["PESA Economic Group Long Name"]],
            "amount": num(r[idx["Amount"]]) or 0.0,
        })
    wb.close()
    return out
